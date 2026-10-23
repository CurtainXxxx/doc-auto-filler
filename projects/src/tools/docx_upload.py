"""Local document persistence and optional standard S3 storage."""
import logging
import os
import re
import uuid
from urllib.parse import urlencode
from docx import Document
from runtime import DATA_DIR
from tools.docx_validator import validate_docx, diff_docx, fix_docx
logger = logging.getLogger(__name__)

def upload_and_validate(docx_bytes: bytes, name_prefix: str, display_name: str, template_path: str, *, validate: bool = True) -> dict:
    """Save locally by default; use boto3 only when S3 is explicitly selected."""
    safe_name = re.sub(r'[^\w.-]', '_', display_name)[:100] or 'document'
    file_name = f"{safe_name}_{uuid.uuid4().hex}.docx"
    output_dir = DATA_DIR / "generated"
    output_dir.mkdir(parents=True, exist_ok=True)
    local_path = output_dir / file_name
    local_path.write_bytes(docx_bytes)
    validation = validate_doc(template_path, str(local_path)) if validate else None
    # Validation can repair XML; publish the same bytes as the local preview.
    docx_bytes = local_path.read_bytes()
    backend = os.getenv("STORAGE_BACKEND", "local").lower()
    if backend == "local":
        download_url = os.getenv("APP_PUBLIC_URL", "").rstrip("/") + "/download-docx?" + urlencode({"file_path": str(local_path)})
    elif backend == "s3":
        import boto3
        bucket = os.getenv("S3_BUCKET_NAME")
        if not bucket:
            raise ValueError("STORAGE_BACKEND=s3 requires S3_BUCKET_NAME")
        client = boto3.client("s3", endpoint_url=os.getenv("S3_ENDPOINT_URL") or None, region_name=os.getenv("S3_REGION", "us-east-1"), aws_access_key_id=os.getenv("S3_ACCESS_KEY_ID") or None, aws_secret_access_key=os.getenv("S3_SECRET_ACCESS_KEY") or None)
        key = name_prefix.strip("/") + "/" + file_name
        client.put_object(Bucket=bucket, Key=key, Body=docx_bytes, ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        download_url = client.generate_presigned_url("get_object", Params={"Bucket": bucket, "Key": key}, ExpiresIn=86400)
    else:
        raise ValueError("STORAGE_BACKEND must be local or s3")
    return {"file_name": file_name, "download_url": download_url, "local_path": str(local_path), "validation": validation}

def validate_doc(template_path: str, output_path: str) -> dict:
    """对比模板和生成文档的结构完整性（增强版校验管线）。

    整合了 docx_validator 模块的多层校验：
    1. 生成文件能否被 python-docx 正常打开
    2. OpenXML元素顺序校验（pPr→runs, rPr→t, tcPr→p, sectPr位置）
    3. 单元格最少段落校验
    4. 表格维度对比（数量、行数、列数）
    5. 合并单元格连续性校验
    6. 格式污染检测
    7. 模板与填写文档对比（diff）

    Args:
        template_path: 原始模板文件路径
        output_path: 生成的文档文件路径

    Returns:
        {"valid": bool, "errors": [...], "warnings": [...], "diff": dict, "checks": list}
    """
    errors: list[str] = []
    warnings: list[str] = []

    # 1. 检查生成文件能否打开
    try:
        output_doc = Document(output_path)
    except Exception as e:
        return {"valid": False, "errors": [f"生成的文档无法打开: {e}"], "warnings": [], "diff": None, "checks": []}

    try:
        template_doc = Document(template_path)
    except Exception as e:
        warnings.append(f"模板文件无法打开，跳过结构对比: {e}")
        return {"valid": True, "errors": [], "warnings": warnings, "diff": None, "checks": []}

    # 2. 运行校验管线（validate_docx）
    validation = validate_docx(output_doc)
    if not validation["passed"]:
        try:
            fix_result = fix_docx(output_path)
            if fix_result["fixed"]:
                warnings.append(f"已自动修复 {len(fix_result['fixes'])} 个结构问题")
                output_doc = Document(output_path)
                validation = validate_docx(output_doc)
        except Exception:
            logger.exception("Document structure repair failed")
            warnings.append("文档结构自动修复失败，请检查校验结果")


    # 合并管线校验结果
    errors.extend(validation["errors"])
    warnings.extend(validation["warnings"])

    # 3. 表格数量和维度对比
    t_tables = len(template_doc.tables)
    o_tables = len(output_doc.tables)
    if t_tables != o_tables:
        errors.append(f"表格数量变化: 模板{t_tables}个, 生成{o_tables}个")

    for i in range(min(t_tables, o_tables)):
        t_rows = len(template_doc.tables[i].rows)
        o_rows = len(output_doc.tables[i].rows)
        if t_rows != o_rows:
            if o_rows < t_rows:
                errors.append(f"表格{i+1}行数减少: 模板{t_rows}行, 生成{o_rows}行")
            else:
                warnings.append(f"表格{i+1}行数增加: 模板{t_rows}行, 生成{o_rows}行（可能是行组填充）")

        t_cols = len(template_doc.tables[i].columns)
        o_cols = len(output_doc.tables[i].columns)
        if t_cols != o_cols:
            errors.append(f"表格{i+1}列数变化: 模板{t_cols}列, 生成{o_cols}列")

    # 4. 生成diff对比
    diff_result = None
    try:
        diff_result = diff_docx(template_path, output_path)
    except Exception as e:
        warnings.append(f"文档对比失败: {e}")

    return {
        "valid": len(errors) == 0,
        "errors": errors,
        "warnings": warnings,
        "diff": diff_result,
        "checks": validation["checks"],
    }
