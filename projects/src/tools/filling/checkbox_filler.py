"""勾选框行检测和填充：选项行标记和 □N. 未勾选复选框检测。

提供两遍扫描策略：
1. 标签+选项+空白格模式：检测"选项 → 空白格"对并填充 √
2. □N. 模式：检测未勾选复选框，根据相邻内容有值时改为 ☑

依赖：filling.base (get_unique_cells, is_vmerge_continue, set_cell_text)
"""

from typing import Dict, List, Optional, Any
from docx.oxml.ns import qn
from tools.filling.base import get_unique_cells, is_vmerge_continue, set_cell_text

# ── 选项词集合 ──
OPTION_WORDS_SET = frozenset({
    "选修", "必修", "开卷", "闭卷", "半开卷", "是", "否",
    "试题库", "试卷库", "教师组题", "本人阅卷", "同行阅卷",
    "集体阅卷", "机器阅卷", "其他",
})


def normalize_option_text(text: str) -> str:
    """标准化选项文字：去除多余空格，使'闭  卷'匹配'闭卷'。"""
    return text.replace(" ", "").replace("　", "")


def find_section_title_above(table, row_idx: int) -> str:
    """从上方行查找章节标题（如'二、考试方式'），用于无行标签的勾选框组。"""
    for r in range(row_idx - 1, max(row_idx - 3, -1), -1):
        row = table.rows[r]
        seen = set()
        for cell in row.cells:
            elem_id = id(cell._element)
            if elem_id not in seen:
                seen.add(elem_id)
                text = cell.text.strip()
                tc = cell._tc
                tcPr = tc.find(qn("w:tcPr"))
                if tcPr is not None:
                    gs = tcPr.find(qn("w:gridSpan"))
                    if gs is not None:
                        span = int(gs.get(qn("w:val"), "1"))
                        if span >= 5 and text:  # 跨5列以上视为章节标题
                            return text
    return ""


def detect_checkbox_row(unique_cells):
    """检测一行是否是勾选框行，返回标签组列表。

    勾选框行模式: [行标签] [选项1] [空白1] [选项2] [空白2] ...
    也支持无行标签但有上方章节标题的模式。

    Returns:
        list[dict]: [{"label": 行标签, "option_blanks": {选项文字: 空白格索引}}, ...]
    """
    groups = []
    current_label = None
    current_options = {}

    for ci in range(len(unique_cells) - 1):
        cell_text = unique_cells[ci].text.strip()
        next_text = unique_cells[ci + 1].text.strip()
        normalized_text = normalize_option_text(cell_text)

        if (normalized_text in OPTION_WORDS_SET or cell_text in OPTION_WORDS_SET) and not next_text:
            opt_key = normalized_text if normalized_text in OPTION_WORDS_SET else cell_text
            current_options[opt_key] = ci + 1
        elif cell_text and normalized_text not in OPTION_WORDS_SET and cell_text not in OPTION_WORDS_SET and not is_vmerge_continue(unique_cells[ci]):
            if current_options and current_label:
                groups.append({"label": current_label, "option_blanks": current_options})
            elif current_options:
                groups.append({"label": current_label or "_checkbox_group", "option_blanks": current_options})
            current_label = cell_text
            current_options = {}

    if current_options:
        if current_label:
            groups.append({"label": current_label, "option_blanks": current_options})
        else:
            groups.append({"label": "_checkbox_group", "option_blanks": current_options})

    return groups


def fill_checkbox_row(unique_cells, group, user_value):
    """根据用户值在勾选框组打√。

    Args:
        unique_cells: 去重后的单元格列表
        group: detect_checkbox_row 返回的标签组
        user_value: 如 "必修" 或 "闭卷" （单选，支持逗号分隔多选）
    """
    selected = [v.strip() for v in user_value.replace("，", ",").split(",")]
    option_blanks = group["option_blanks"]

    # 找到同行中第一个有格式的cell作为格式源
    rPr_source = None
    for cell in unique_cells:
        tc = cell._element
        for p in tc.findall(qn("w:p")):
            for r in p.findall(qn("w:r")):
                rPr = r.find(qn("w:rPr"))
                if rPr is not None:
                    rPr_source = rPr
                    break
            if rPr_source is not None:
                break
        if rPr_source is not None:
            break

    for opt_text, blank_idx in option_blanks.items():
        if opt_text in selected:
            set_cell_text(unique_cells[blank_idx], "√", rPr_source=rPr_source)


def fill_checkbox_rows_in_table(doc, analysis, data):
    """扫描所有表格行，检测勾选框行并填充。

    两遍扫描：
    1. 标签+选项+空白格模式
    2. □N. 未勾选复选框模式（根据相邻内容字段在 data 中是否有值来判断是否勾选）
    """
    # 第一遍：标签+选项+空白格模式
    for t_idx, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            unique = get_unique_cells(row)
            groups = detect_checkbox_row(unique)

            if not groups:
                continue

            for group in groups:
                row_label = group["label"]
                user_value = None

                # 1. 精确匹配
                if row_label in data:
                    user_value = data[row_label]
                else:
                    # 2. 模糊匹配
                    clean_label = row_label.replace("\n", "").replace(" ", "")
                    for key in data:
                        clean_key = key.replace("\n", "").replace(" ", "")
                        if clean_key == clean_label:
                            user_value = data[key]
                            break

                # 3. 对_checkbox_group标签（无行标签），尝试从上方章节标题匹配
                if user_value is None and row_label == "_checkbox_group":
                    section_title = find_section_title_above(table, ri)
                    if section_title:
                        clean_section = section_title.replace("\n", "").replace(" ", "")
                        for key in data:
                            clean_key = key.replace("\n", "").replace(" ", "")
                            if clean_key == clean_section or clean_key in clean_section:
                                user_value = data[key]
                                break

                # 4. 检查选项本身是否在data中
                if user_value is None:
                    for opt in group["option_blanks"]:
                        if opt in data:
                            user_value = data[opt]
                            break

                # 5. 对_checkbox_group，检查data中是否有该组任意选项的同义词
                if user_value is None and row_label == "_checkbox_group":
                    for opt in group["option_blanks"]:
                        for key in data:
                            if normalize_option_text(key) == normalize_option_text(opt):
                                user_value = data[key]
                                break
                        if user_value is not None:
                            break

                if user_value is not None:
                    fill_checkbox_row(unique, group, user_value)

    # 第二遍：检测 □N. 模式（未勾选的复选框），根据相邻内容字段在 data 中是否有值来判断是否勾选
    for t_idx, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            unique = get_unique_cells(row)
            for ci, cell in enumerate(unique):
                text = cell.text.strip()
                if text.startswith('□') and len(text) > 1:
                    has_value = False
                    for cj in range(ci + 1, len(unique)):
                        neighbor_text = unique[cj].text.strip()
                        for sep in ['：', ':']:
                            if sep in neighbor_text:
                                label_candidate = neighbor_text.split(sep)[0].strip()
                                if not label_candidate:
                                    continue
                                if label_candidate in data and data[label_candidate]:
                                    has_value = True
                                    break
                                clean_label = label_candidate.replace("\n", "").replace(" ", "")
                                for key in data:
                                    clean_key = key.replace("\n", "").replace(" ", "")
                                    if clean_key == clean_label or clean_label in clean_key or clean_key in clean_label:
                                        if data[key]:
                                            has_value = True
                                            break
                                if has_value:
                                    break
                        if has_value:
                            break
                    if has_value:
                        for p in cell.paragraphs:
                            for run in p.runs:
                                if '□' in run.text:
                                    run.text = run.text.replace('□', '☑', 1)
                                    break
                            else:
                                continue
                            break