"""Regression coverage for replacing the hosting runtime and storage."""
import asyncio
import io
import json
from pathlib import Path
from unittest.mock import MagicMock
import pytest
from docx import Document
from fastapi.testclient import TestClient
from langchain_core.messages import AIMessage, HumanMessage
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.outputs import ChatResult, ChatGeneration
from runtime import PROJECT_ROOT, Context, request_context
from main import app, service, _safe_path

client = TestClient(app)

def template_bytes():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(1, 0).text = "联系电话"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()

def test_no_platform_packages_or_imports():
    import importlib.metadata
    assert not any(d.metadata["Name"].lower().startswith("coze") for d in importlib.metadata.distributions())
    for path in (PROJECT_ROOT / "src").rglob("*.py"):
        source = path.read_text(encoding="utf-8")
        assert "coze_coding" not in source
        assert "coze_workload" not in source
        assert "COZE_" not in source

def test_health_and_frontend_without_model():
    assert client.get("/health").json()["status"] == "ok"
    assert client.get("/web/").status_code == 200

def test_upload_preview_generate_and_download():
    response = client.post("/upload-template", files={"file": ("form.docx", template_bytes())})
    path = response.json()["template_path"]
    preview = client.get("/template-preview", params={"path": path}).json()
    assert preview["success"]
    assert preview["field_map"]
    result = client.post("/cofill-generate", json={"template_path": path, "field_values": {"姓名": "张三", "联系电话": "13800138000"}}).json()
    assert result["success"], result
    download = client.get(result["download_url"])
    assert download.status_code == 200
    doc = Document(io.BytesIO(download.content))
    text = " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "张三" in text and "13800138000" in text
    assert client.get("/generated-preview", params={"local_path": result["local_path"]}).json()["success"]

def test_knowledge_upload_without_model():
    result = client.post("/upload", files={"files": ("facts.txt", "姓名：张三".encode())}).json()
    assert result["success"]
    assert Path(result["file_path"]).is_file()

@pytest.mark.parametrize("name", ["评价报告", "试卷分析", "关联矩阵"])
def test_builtin_preview(name):
    from tools.edu_report_tool import _get_template_path
    from tools.old_report_extractor import _resolve_template_path
    assert Path(_get_template_path(name)).is_file()
    assert Path(_resolve_template_path(name)).is_file()
    result = client.get("/template-preview", params={"path": name}).json()
    assert result["success"], result
    field = next(iter(result["field_map"].values()))
    generated = client.post("/cofill-generate", json={
        "template_path": name, "field_values": {field["label"]: "测试值"}
    }).json()
    assert generated["success"], generated
    assert client.get(generated["download_url"]).status_code == 200

def test_path_boundary():
    from fastapi import HTTPException
    with pytest.raises(HTTPException):
        _safe_path(str(PROJECT_ROOT.parent / (PROJECT_ROOT.name + "-other") / "secret"))

def test_model_config_shared_and_explicit(monkeypatch):
    from model_provider import build_model
    for key in ["EXTERNAL_LLM_API_KEY", "EXTERNAL_LLM_BASE_URL", "EXTERNAL_LLM_MODEL"]:
        monkeypatch.delenv(key, raising=False)
    with pytest.raises(ValueError, match="EXTERNAL_LLM"):
        build_model()
    monkeypatch.setenv("EXTERNAL_LLM_API_KEY", "offline-test")
    monkeypatch.setenv("EXTERNAL_LLM_BASE_URL", "http://localhost:12345/v1")
    monkeypatch.setenv("EXTERNAL_LLM_MODEL", "local-test")
    assert build_model().model_name == "local-test"

def test_legacy_message_input():
    from service import GraphService
    payload = GraphService()._input({"message": "分析模板", "session_id": "legacy"})
    assert payload["messages"][0]["content"] == "分析模板"

def test_s3_uses_standard_client(monkeypatch):
    import boto3
    from tools.docx_upload import upload_and_validate
    mock = MagicMock()
    mock.generate_presigned_url.return_value = "https://storage.example/report.docx"
    monkeypatch.setattr(boto3, "client", lambda *a, **kw: mock)
    monkeypatch.setenv("STORAGE_BACKEND", "s3")
    monkeypatch.setenv("S3_BUCKET_NAME", "reports")
    result = upload_and_validate(template_bytes(), "reports", "sample", "", validate=False)
    assert result["download_url"] == "https://storage.example/report.docx"
    assert mock.put_object.call_args.kwargs["Bucket"] == "reports"
    monkeypatch.delenv("S3_BUCKET_NAME")
    with pytest.raises(ValueError, match="S3_BUCKET_NAME"):
        upload_and_validate(template_bytes(), "reports", "sample", "", validate=False)

class OfflineModel(BaseChatModel):
    """Exercise real nested agents and tool execution without a network call."""
    @property
    def _llm_type(self):
        return "offline-test"

    def bind_tools(self, tools, **kwargs):
        return self

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        if not any(getattr(m, "type", "") == "tool" for m in messages):
            msg = AIMessage(content="", tool_calls=[{"id": "test-call", "name": "get_fill_checklist", "args": {"template_name_or_path": "评价报告"}}])
        else:
            msg = AIMessage(content="请补充需要填写的信息")
        return ChatResult(generations=[ChatGeneration(message=msg)])

def test_real_graph_sessions_and_tool_events(monkeypatch):
    import agents.agent as agent
    from service import GraphService
    monkeypatch.setattr(agent, "_build_llm", lambda ctx=None: OfflineModel())
    graph_service = GraphService()
    async def run():
        a = await graph_service.run({"session_id": "test-a", "text": "你好"}, Context())
        b = await graph_service.run({"session_id": "test-b", "text": "独立会话"}, Context())
        assert not any(isinstance(m, HumanMessage) and m.content == "你好" for m in b["messages"])
        a2 = await graph_service.run({"session_id": "test-a", "text": "继续"}, Context())
        assert any(isinstance(m, HumanMessage) and m.content == "你好" for m in a2["messages"])
        events = [e async for e in graph_service.deltas({"session_id": "test-c", "text": "你好"}, Context())]
        assert any(e.get("role") == "tool" for e in events), events
        assert any(e.get("content") == "请补充需要填写的信息" for e in events), events
        assert request_context.get() is None
    asyncio.run(run())

def test_http_chat_stream_and_nonstream(monkeypatch):
    import agents.agent as agent
    monkeypatch.setattr(agent, "_build_llm", lambda ctx=None: OfflineModel())
    monkeypatch.setattr(service, "_graph", None)
    response = client.post("/v1/chat/completions", json={"messages": [{"role": "user", "content": "你好"}], "session_id": "http-sync"})
    assert response.status_code == 200, response.text
    assert response.json()["choices"][0]["message"]["content"]
    response = client.post("/v1/chat/completions", json={"messages": [{"role": "user", "content": "你好"}], "stream": True, "session_id": "http-stream"})
    assert "data: [DONE]" in response.text
    chunks = [json.loads(line[6:]) for line in response.text.splitlines() if line.startswith("data: {")]
    assert not any("error" in c for c in chunks), chunks
    assert any(c["choices"][0]["delta"].get("role") == "tool" for c in chunks)

def test_cancel_cleans_up_producer(monkeypatch):
    from service import GraphService
    svc = GraphService()
    async def slow(payload, ctx):
        await asyncio.sleep(100)
        yield {"content": "late"}
    monkeypatch.setattr(svc, "deltas", slow)
    async def run():
        ctx = Context()
        stream = svc.stream_sse({}, ctx)
        consumer = asyncio.create_task(anext(stream))
        await asyncio.sleep(0.02)
        assert svc.cancel_run(ctx.run_id)["status"] == "success"
        consumer.cancel()
        await asyncio.gather(consumer, return_exceptions=True)
        await stream.aclose()
        assert not svc.running_tasks
    asyncio.run(run())
