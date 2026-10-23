"""Behavioral regressions for the shared legacy Word filling pipeline."""
import copy
import io

import pytest
from docx import Document

from tools.filling.builtin_expander import expand_report_data
from tools.filling.doc_builder import build_report_docx, fill_custom_template


def test_builtin_expansion_does_not_mutate_caller_data():
    analysis = {"label_fields": [{"label": "课程_应到", "field_id": "A"}]}
    values = {"应到": 0, "其他": False}
    original = copy.deepcopy(values)
    assert expand_report_data("unused", values, analysis)["课程_应到"] == "0"
    assert values == original


def test_multi_column_values_use_field_id_before_label(tmp_path):
    doc = Document()
    table = doc.add_table(rows=1, cols=3)
    table.cell(0, 0).text = "课程目标"
    path = tmp_path / "multi.docx"
    doc.save(path)
    field = {"field_id": "T0_MULTI", "label": "课程目标", "pattern": "multi_col",
             "fill_mode": "group", "table_idx": 0, "row_idx": 0, "row_indices": [0],
             "col_idx": 1, "fillable_cols": [1, 2]}
    analysis = {"label_fields": [field], "row_groups": []}
    result, *_ = fill_custom_template(str(path), {"T0_MULTI": [0, False], "课程目标": "wrong,wrong"}, analysis)
    output = Document(io.BytesIO(result))
    assert [c.text for c in output.tables[0].rows[0].cells] == ["课程目标", "0", "False"]


@pytest.mark.parametrize("builder", [build_report_docx, fill_custom_template])
def test_missing_values_do_not_write_literal_none(tmp_path, builder):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(1, 0).text = "成绩"
    path = tmp_path / "blank.docx"
    doc.save(path)
    fields = [{"field_id": f"F{i}", "label": label, "pattern": "label_blank", "fill_mode": "set",
               "table_idx": 0, "row_idx": i, "row_indices": [i], "col_idx": 1}
              for i, label in enumerate(["姓名", "成绩"])]
    output, *_ = builder(str(path), {"F0": None, "F1": 0}, {"label_fields": fields, "row_groups": []})
    result = Document(io.BytesIO(output))
    assert result.tables[0].cell(0, 1).text == ""
    assert result.tables[0].cell(1, 1).text == "0"
