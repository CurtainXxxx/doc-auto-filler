"""Regressions for confirmed review findings; no external model or cloud service."""
import asyncio
import copy
import io
import json
from pathlib import Path
import time
import zipfile
import pytest
from docx import Document
from fastapi.testclient import TestClient
from main import app
from document_access import MAX_FILE_BYTES, MAX_REQUEST_BYTES, validate_download_url, check_archive
from runtime import Context
from runtime import PROJECT_ROOT, DATA_DIR
from tools.form_filling_state import FormFillingState
from tools.filling.row_group_filler import fill_simple_row_groups
from tools.filling.paragraph_filler import fill_paragraph_fields
from tools.prefill_tool import _merge_multi_file_results
from tools.docx_preview import _render_paragraph_with_fields

client = TestClient(app)


@pytest.mark.parametrize("template,values,expected", [
    ("考场记录表", {"考试方式": "闭卷"}, {6: {"闭卷"}}),
    ("考场记录表", {"T0_R6_CHK": "开卷", "考试方式": "闭卷"}, {6: {"开卷"}}),
    ("试卷分析模板", {"课程性质": "必修", "命题形式": "教师组题", "考试形式": "闭卷", "教考分离": "否"},
     {3: {"必修"}, 4: {"教师组题"}, 5: {"闭卷", "否"}}),
])
def test_manual_http_generation_selects_the_correct_checkbox(template, values, expected):
    from tools.filling.base import get_unique_cells
    from tools.filling.checkbox_filler import detect_checkbox_row
    template_path = PROJECT_ROOT / "assets" / "templates" / f"{template}.docx"
    result = client.post("/cofill-generate", json={"template_path": str(template_path), "field_values": values}).json()
    assert result["success"], result
    doc = Document(io.BytesIO(client.get(result["download_url"]).content))
    # Detect positions from the pristine template, since selected blanks contain √.
    source = Document(template_path)
    for row_idx, selected in expected.items():
        source_cells = get_unique_cells(source.tables[0].rows[row_idx])
        output_cells = get_unique_cells(doc.tables[0].rows[row_idx])
        for group in detect_checkbox_row(source_cells):
            for option, blank_idx in group["option_blanks"].items():
                assert output_cells[blank_idx].text.strip() == ("√" if option in selected else "")


def upload_doc(doc):
    buffer = io.BytesIO()
    doc.save(buffer)
    result = client.post("/upload-template", files={"file": ("form.docx", buffer.getvalue())})
    assert result.status_code == 200, result.text
    return result.json()["template_path"]


def test_zero_false_and_session_roundtrip():
    state = FormFillingState(session_id="kept", analysis_result={"label_fields": [
        {"field_id": "a", "label": "分数"}, {"field_id": "b", "label": "通过"}]})
    state.bulk_fill({"a": 0, "b": False})
    assert state.get_field_values() == {"a": "0", "b": "False"}
    state.confirm_fields([])
    assert state.get_progress()["confirmed"] == 0
    assert FormFillingState.from_json(state.to_json()).session_id == "kept"
    assert state._resolve_field_id("") is None
    state.bulk_fill({"T0_G0": [[1]]})
    state.init_from_analysis({"label_fields": []})
    assert state.get_label_value_map() == {}


def test_row_expansion_writes_new_values_in_order():
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "课程"
    table.cell(0, 1).text = "分数"
    fill_simple_row_groups(doc, [{"group_id": "T0_G0", "table_idx": 0, "start_row": 1, "template_row_count": 1}],
                           {"T0_G0": [["A", 1], ["B", 2], ["C", 3]]})
    assert [[c.text for c in r.cells] for r in table.rows][1:] == [["A", "1"], ["B", "2"], ["C", "3"]]


def test_expanding_one_group_keeps_following_group_coordinates():
    doc = Document()
    table = doc.add_table(rows=4, cols=1)
    table.cell(0, 0).text = "第一组"
    table.cell(2, 0).text = "第二组"
    groups = [{"group_id": "T0_G0", "table_idx": 0, "start_row": 1, "template_row_count": 1},
              {"group_id": "T0_G1", "table_idx": 0, "start_row": 3, "template_row_count": 1}]
    fill_simple_row_groups(doc, groups, {"T0_G0": [["A"], ["B"]], "T0_G1": [["C"], ["D"]]})
    assert [r.cells[0].text for r in table.rows] == ["第一组", "A", "B", "第二组", "C", "D"]


def test_paragraph_values_are_literal_and_ids_take_priority():
    doc = Document()
    doc.add_paragraph("路径：____")
    fill_paragraph_fields(doc, [{"label": "路径", "field_id": "P0", "row_idx": 0}],
                          {"路径": "wrong", "P0": r"C:\new\docs"})
    assert r"C:\new\docs" in doc.paragraphs[0].text


def test_multifile_merge_preserves_inputs_order_and_structured_values():
    source = [[{"field_id": "b", "label": "b", "value": [1], "confidence": 1., "source": "a"},
               {"field_id": "a", "label": "a", "value": "A", "confidence": .8, "source": "a"}],
              [{"field_id": "b", "label": "b", "value": [2], "confidence": .9, "source": "b"}]]
    before = copy.deepcopy(source)
    result = _merge_multi_file_results(source)
    assert source == before
    assert [r["field_id"] for r in result] == ["b", "a"]
    assert result[0]["confidence"] == .6


def test_preview_escapes_original_paragraph():
    rendered = _render_paragraph_with_fields('<img src=x onerror="alert(1)"> 姓名：____', [("P0", {"label": "姓名"})], "")
    assert "<img" not in rendered
    assert "&lt;img" in rendered
    assert '<span class="field-blank' in rendered


@pytest.mark.parametrize("path", [".env", "src/main.py", "config/agent_llm_config.json"])
def test_download_cannot_read_source_or_config(path):
    assert client.get("/download-docx", params={"file_path": path}).status_code == 403


def test_global_temp_is_not_a_document_directory(tmp_path):
    target = tmp_path / "unrelated.docx"
    target.write_bytes(b"unrelated")
    assert client.get("/download-docx", params={"file_path": str(target)}).status_code == 403


def test_remote_download_denied_before_network_call(monkeypatch):
    monkeypatch.delenv("DOWNLOAD_ALLOWED_ORIGINS", raising=False)
    monkeypatch.delenv("S3_ENDPOINT_URL", raising=False)
    assert client.get("/download-docx", params={"file_path": "http://127.0.0.1/internal"}).status_code == 403
    monkeypatch.setenv("DOWNLOAD_ALLOWED_ORIGINS", "https://files.example")
    assert validate_download_url("https://files.example/report.docx")
    for url in ["https://files.example.evil/report.docx", "https://user:password@files.example/report.docx"]:
        with pytest.raises(PermissionError):
            validate_download_url(url)


def test_upload_limits_and_invalid_word():
    assert client.post("/upload-template", content=b"", headers={"Content-Length": str(MAX_REQUEST_BYTES + 1)}).status_code == 413
    assert client.post("/upload-template", files={"file": ("large.docx", b"a" * (MAX_FILE_BYTES + 1))}).status_code == 413
    assert client.post("/upload-template", files={"file": ("invalid.docx", b"invalid")}).status_code == 422
    invalid_knowledge = client.post("/upload", files={"files": ("invalid.docx", b"invalid")})
    assert invalid_knowledge.status_code == 200
    assert not invalid_knowledge.json()["success"]
    assert client.post("/upload", files=[("files", (f"{i}.txt", b"x")) for i in range(11)]).status_code == 413


def test_archive_expanded_size_is_checked(tmp_path):
    target = tmp_path / "too-many.docx"
    with zipfile.ZipFile(target, "w") as archive:
        for i in range(2049):
            archive.writestr(str(i), b"")
    with pytest.raises(ValueError):
        check_archive(str(target))


@pytest.mark.parametrize("endpoint", ["/run", "/v1/chat/completions", "/prefill"])
def test_nonobject_json_returns_client_error(endpoint):
    assert client.post(endpoint, json=[]).status_code == 400


def test_http_zero_value_fills_document():
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "成绩"
    path = upload_doc(doc)
    result = client.post("/cofill-generate", json={"template_path": path, "field_values": {"成绩": 0}}).json()
    assert result["success"], result
    output = Document(io.BytesIO(client.get(result["download_url"]).content))
    assert "0" in " ".join(c.text for t in output.tables for r in t.rows for c in r.cells)


def test_form_generation_preserves_duplicate_label_values():
    from tools.edu_report_tool import _active_form_states, generate_form_document
    doc = Document()
    doc.add_paragraph("日期：____")
    doc.add_paragraph("日期：____")
    path = upload_doc(doc)
    fields = [{"field_id": "P0", "label": "日期", "raw_label": "日期", "fill_mode": "append", "pattern": "paragraph_underline", "table_idx": -1,
               "row_idx": 0, "row_indices": [0], "col_idx": -1},
              {"field_id": "P1", "label": "日期", "raw_label": "日期", "fill_mode": "append", "pattern": "paragraph_underline", "table_idx": -1,
               "row_idx": 1, "row_indices": [1], "col_idx": -1}]
    state = FormFillingState(session_id="duplicate-labels", template_path=path, analysis_result={"label_fields": fields})
    state.bulk_fill({"P0": "2026-01-01", "P1": "2026-02-02"})
    _active_form_states[state.session_id] = state
    from tools.edu_report_tool import _cache_analysis
    _cache_analysis(path, {"label_fields": fields, "row_groups": []})
    try:
        result = json.loads(generate_form_document(session_id=state.session_id))
        assert result["success"], result
        output = Document(result["local_path"])
        assert "2026-01-01" in output.paragraphs[0].text
        assert "2026-02-02" in output.paragraphs[1].text
    finally:
        _active_form_states.pop(state.session_id, None)


def test_prefill_work_does_not_block_event_loop(monkeypatch):
    import httpx
    import tools.prefill_tool as prefill
    def slow(*args):
        time.sleep(.2)
        return {"success": True}
    monkeypatch.setattr(prefill, "prefill_from_file_paths", slow)
    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as http:
            pending = asyncio.create_task(http.post("/prefill", json={"template_path": "评价报告", "file_paths": ["assets/knowledge/test.txt"]}))
            await asyncio.sleep(.02)
            assert not pending.done()
            assert (await http.get("/health")).status_code == 200
            assert (await pending).json()["success"]
    asyncio.run(run())


def test_same_session_concurrency_is_rejected_and_released(monkeypatch):
    from service import GraphService, SessionBusyError
    svc = GraphService()
    started = asyncio.Event()
    release = asyncio.Event()
    class SlowGraph:
        async def ainvoke(self, *args, **kwargs):
            started.set()
            await release.wait()
            return {"messages": []}
    monkeypatch.setattr(svc, "_get_graph", lambda ctx: SlowGraph())
    async def run():
        task = asyncio.create_task(svc.run({"session_id": "same", "text": "a"}, Context()))
        await started.wait()
        with pytest.raises(SessionBusyError):
            await svc.run({"session_id": "same", "text": "b"}, Context())
        release.set()
        await task
        assert not svc._active_sessions
    asyncio.run(run())


def test_stream_error_does_not_claim_success(monkeypatch):
    from service import GraphService
    svc = GraphService()
    async def fail(*args):
        raise ValueError("test failure")
        yield  # Make this a failing asynchronous generator.
    monkeypatch.setattr(svc, "deltas", fail)
    async def run():
        output = "".join([x async for x in svc.stream_sse({}, Context(), openai=True)])
        assert '"error"' in output
        assert '"finish_reason": "stop"' not in output
        assert "data: [DONE]" in output
        assert not svc.running_tasks
    asyncio.run(run())


@pytest.mark.parametrize("template", sorted((PROJECT_ROOT / "assets" / "templates").glob("*.docx")), ids=lambda p: p.stem)
def test_shipped_templates_generate_and_remain_unchanged(template):
    from tools.template_analyzer import analyze_template
    from tools.filling.doc_builder import fill_custom_template
    before = template.read_bytes()
    analysis = analyze_template(str(template))
    values = {field["field_id"]: "review value" for field in analysis["label_fields"]
              if field.get("fill_mode") != "group"}
    output, _, _, _ = fill_custom_template(str(template), values, analysis)
    generated = Document(io.BytesIO(output))
    assert len(generated.tables) == len(Document(str(template)).tables)
    assert zipfile.is_zipfile(io.BytesIO(output))
    assert template.read_bytes() == before


def test_remote_redirect_and_response_size_are_limited(monkeypatch):
    import httpx
    import main
    original_client = httpx.AsyncClient
    monkeypatch.setenv("DOWNLOAD_ALLOWED_ORIGINS", "https://files.example")
    requests = []
    def redirect(request):
        requests.append(str(request.url))
        return httpx.Response(302, headers={"Location": "http://127.0.0.1/private"})
    def fake_client(*args, **kwargs):
        return original_client(transport=httpx.MockTransport(redirect), **kwargs)
    monkeypatch.setattr(httpx, "AsyncClient", fake_client)
    result = client.get("/download-docx", params={"file_path": "https://files.example/report.docx"}).json()
    assert not result["success"]
    assert requests == ["https://files.example/report.docx"]
    before = set(DATA_DIR.glob("tmp*.docx"))
    monkeypatch.setattr(main, "MAX_FILE_BYTES", 5)
    monkeypatch.setattr(httpx, "AsyncClient", lambda *args, **kwargs: original_client(
        transport=httpx.MockTransport(lambda request: httpx.Response(200, content=b"123456")), **kwargs))
    assert not client.get("/download-docx", params={"file_path": "https://files.example/report.docx"}).json()["success"]
    assert set(DATA_DIR.glob("tmp*.docx")) == before


def test_s3_publishes_validated_bytes(monkeypatch):
    from unittest.mock import MagicMock
    import boto3
    import tools.docx_upload as storage
    mock = MagicMock()
    mock.generate_presigned_url.return_value = "https://files.example/result.docx"
    monkeypatch.setattr(boto3, "client", lambda *args, **kwargs: mock)
    monkeypatch.setenv("STORAGE_BACKEND", "s3")
    monkeypatch.setenv("S3_BUCKET_NAME", "reports")
    def repair(template, path):
        Path(path).write_bytes(b"repaired")
        return {"valid": True}
    monkeypatch.setattr(storage, "validate_doc", repair)
    result = storage.upload_and_validate(b"original", "reports", "sample", "unused")
    assert mock.put_object.call_args.kwargs["Body"] == Path(result["local_path"]).read_bytes() == b"repaired"


def test_chunked_request_limit_without_content_length():
    import httpx
    from fastapi import FastAPI, Request
    from http_limits import RequestSizeLimit
    small = FastAPI()
    small.add_middleware(RequestSizeLimit, limit=10)
    @small.post("/")
    async def consume(request: Request):
        await request.body()
        return {"ok": True}
    async def chunks():
        yield b"123456"
        yield b"123456"
    async def run():
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=small), base_url="http://test") as http:
            response = await http.post("/", content=chunks())
            assert response.status_code == 413
    asyncio.run(run())


def test_old_report_prefill_uses_the_active_form_session():
    import tools.edu_report_tool as forms
    import tools.old_report_extractor as reports
    from agents.agent import _load_tools
    _load_tools()
    template = Document()
    table = template.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(1, 0).text = "联系电话"
    template_path = upload_doc(template)
    old = Document(template_path)
    old.tables[0].cell(0, 1).text = "张三"
    old.tables[0].cell(1, 1).text = "13800138000"
    old_path = upload_doc(old)
    result = forms.init_form_filling.invoke({"session_id": "old-report-review", "template_name_or_path": template_path})
    assert json.loads(result)["success"]
    try:
        result = json.loads(reports.prefill_from_old_report.invoke({"session_id": "old-report-review", "old_report_path": old_path}))
        assert result["success"], result
        assert result["prefilled_count"] > 0, result
        assert result["prefilled_fields"], result
        assert "张三" in forms._active_form_states["old-report-review"].get_field_values().values()
    finally:
        forms._active_form_states.pop("old-report-review", None)
