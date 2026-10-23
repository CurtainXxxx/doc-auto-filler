"""Drive the real supervisor graph and document engine without paid model calls."""
import asyncio
import copy
import io
import json
import uuid
from pathlib import Path

import pytest
from docx import Document
from langchain_core.language_models.chat_models import BaseChatModel
from langchain_core.messages import AIMessage, ToolMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from langgraph.checkpoint.memory import MemorySaver
from pydantic import Field

from agents.document_workers import DocumentSession
from agents.leader_worker import MAX_DECISIONS, build_leader_agent
from document_access import save_upload
from runtime import Context
from service import GraphService


@pytest.mark.parametrize("value,quote,supported", [("0", "成绩：100", False),
    ("10", "成绩：100", False), ("0", "成绩：0", True), ("PB12", "学号：PB123", False)])
def test_numeric_or_identifier_substrings_are_not_evidence(value, quote, supported):
    from agents.document_workers import quoted_value
    assert quoted_value(value, quote) is supported


class SmokeModel(BaseChatModel):
    """Script decisions, while retaining real graph state, workers and DOCX I/O."""
    template_path: str
    file_paths: list[str]
    retry_once: bool = False
    loop: bool = False
    omit_candidate: bool = False
    worker_failed: bool = False
    calls: list[str] = Field(default_factory=list)

    @property
    def _llm_type(self):
        return "leader-smoke-script"

    def bind_tools(self, tools, **kwargs):
        return self

    def _generate(self, messages, stop=None, run_manager=None, **kwargs):
        system = messages[0].content
        if system.startswith("你是材料提取Worker"):
            self.calls.append("extract_worker")
            if self.retry_once and not self.worker_failed:
                self.worker_failed = True
                raise ValueError("simulated transient failure")
            payload = json.loads(messages[-1].content)
            facts = []
            for line in payload["text"].splitlines():
                if "：" in line:
                    label, value = line.split("：", 1)
                    facts.append({"label": label, "value": value, "quote": line})
            reply = AIMessage(content=json.dumps({"facts": facts}, ensure_ascii=False))
        elif system.startswith("你是字段匹配Worker"):
            self.calls.append("matching_worker")
            payload = json.loads(messages[-1].content)
            candidates = [{"field_id": field["field_id"], "value": fact["value"],
                           "evidence_ids": [fact["id"]], "confidence": 1.0}
                          for field in payload["fields"] for fact in payload["facts"]
                          if field["label"] == fact["label"]]
            if self.omit_candidate:
                candidates = candidates[:1]
            reply = AIMessage(content=json.dumps({"candidates": candidates}, ensure_ascii=False))
        else:
            self.calls.append("leader")
            user = next(m.content for m in reversed(messages) if m.type == "human")
            last_user = max(i for i, m in enumerate(messages) if m.type == "human")
            tools = [m for m in messages[last_user + 1:] if isinstance(m, ToolMessage)]
            names = [m.name for m in tools]
            name, args = None, {}
            if self.loop:
                name = "document_status"
            elif "状态" in user:
                if not tools:
                    name = "document_status"
            elif "改为" in user:
                if not tools:
                    name, args = "apply_user_values", {"values": {"T0_R0_C1": "李四"}}
                elif names[-1] == "apply_user_values":
                    name = "generate_document"
            elif tools and names[-1] == "extract_documents" and not json.loads(tools[-1].content)["success"]:
                name, args = "extract_documents", {"file_paths": self.file_paths}
            elif tools and names[-1] == "match_fields" and not json.loads(tools[-1].content)["success"]:
                reply = AIMessage(content="材料存在冲突或遗漏，请确认后继续。")
                return ChatResult(generations=[ChatGeneration(message=reply)])
            else:
                plan = [("inspect_template", {"template_path": self.template_path}),
                        ("extract_documents", {"file_paths": self.file_paths}),
                        ("match_fields", {}), ("commit_proposal", {}),
                        ("validate_document", {}), ("generate_document", {})]
                done = set(names)
                name, args = next(((name, args) for name, args in plan if name not in done), (None, {}))
            reply = AIMessage(content="完成" if name is None else "", tool_calls=[] if name is None else [
                {"id": uuid.uuid4().hex, "name": name, "args": args}])
        reply.usage_metadata = {"input_tokens": 30, "output_tokens": 10, "total_tokens": 40}
        return ChatResult(generations=[ChatGeneration(message=reply)])


def fixture_files(second=None):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(1, 0).text = "成绩"
    buffer = io.BytesIO()
    doc.save(buffer)
    template = save_upload(buffer.getvalue(), ".docx")
    sources = [save_upload("姓名：张三\n成绩：0".encode(), ".txt")]
    if second:
        sources.append(save_upload(second.encode(), ".txt"))
    return template, sources


def service_for(model):
    service = GraphService(mode="leader_worker")
    service._graph = build_leader_agent(model=model, checkpointer=MemorySaver())
    return service


def test_supervisor_real_document_and_targeted_followup():
    template, paths = fixture_files()
    model = SmokeModel(template_path=template, file_paths=paths)
    service = service_for(model)
    async def run():
        result = await service.run({"text": "请按材料填写并生成", "session_id": "smoke"})
        assert result["delivered"] and result["document"]["values"]["T0_R1_C1"] == "0"
        output = Document(result["document"]["artifact"]["local_path"])
        assert output.tables[0].cell(0, 1).text == "张三"
        assert output.tables[0].cell(1, 1).text == "0"
        assert "[生成完成]" in result["messages"][-1].content
        before_calls = len(model.calls)
        changed = await service.run({"text": "姓名改为李四，请重新生成", "session_id": "smoke"})
        assert changed["delivered"]
        assert [item["action"] for item in changed["trace"]] == ["apply_user_values", "generate_document"]
        assert not any("worker" in name for name in model.calls[before_calls:])
        assert Document(changed["document"]["artifact"]["local_path"]).tables[0].cell(0, 1).text == "李四"
        other = await service.run({"text": "查看状态", "session_id": "other"})
        assert not other.get("document", {}).get("values")
        assert not any(m.type == "human" and m.content == "请按材料填写并生成" for m in other["messages"])
    asyncio.run(run())


def test_conflicting_sources_block_generation_and_leave_draft():
    template, paths = fixture_files("姓名：李四\n成绩：0")
    service = service_for(SmokeModel(template_path=template, file_paths=paths))
    result = asyncio.run(service.run({"text": "请根据两份材料生成", "session_id": "conflict"}))
    assert result["document"]["proposal"]["conflicts"]
    assert "artifact" not in result["document"]
    assert not result["delivered"] and "[生成完成]" not in result["messages"][-1].content


def test_extracted_values_omitted_by_matching_do_not_silently_generate():
    template, paths = fixture_files()
    service = service_for(SmokeModel(template_path=template, file_paths=paths, omit_candidate=True))
    result = asyncio.run(service.run({"text": "按材料生成", "session_id": "omission"}))
    assert any("候选遗漏" in error for error in result["document"]["proposal"]["errors"])
    assert not result["delivered"]


def test_supervisor_retries_failed_worker_before_generating():
    template, paths = fixture_files()
    service = service_for(SmokeModel(template_path=template, file_paths=paths, retry_once=True))
    result = asyncio.run(service.run({"text": "按材料生成", "session_id": "retry"}))
    attempts = [t for t in result["trace"] if t["action"] == "extract_documents"]
    assert [t["success"] for t in attempts] == [False, True]
    assert result["delivered"] and not result["document"]["pending_sources"]
    assert not result["metrics"]["usage_complete"]


def test_tool_loop_stops_and_preserves_checkpoint():
    template, paths = fixture_files()
    service = service_for(SmokeModel(template_path=template, file_paths=paths, loop=True))
    result = asyncio.run(service.run({"text": "重复调度", "session_id": "loop"}))
    assert result["decisions"] == MAX_DECISIONS and result["budget_exhausted"]
    assert not result["delivered"] and "预算" in result["messages"][-1].content


def test_supervisor_sse_hides_worker_json_and_emits_verified_result():
    template, paths = fixture_files()
    service = service_for(SmokeModel(template_path=template, file_paths=paths))
    async def run():
        events = [e async for e in service.deltas({"text": "按材料生成", "session_id": "sse"}, Context())]
        assert any(e.get("role") == "tool" and e["name"] == "generate_document" for e in events)
        text = "".join(e.get("content", "") for e in events if e.get("role") != "tool")
        assert "[生成完成]" in text and '"facts"' not in text
    asyncio.run(run())


def test_generation_checks_output_before_publication(monkeypatch):
    import agents.document_workers as workers
    template, paths = fixture_files()
    model = SmokeModel(template_path=template, file_paths=paths)
    async def run():
        session = DocumentSession({}, model, {}, "按材料生成", {"model_calls": 0, "usage_complete": True})
        await session.inspect_template(template)
        await session.extract_documents(paths)
        await session.match_fields()
        session.commit_proposal()
        monkeypatch.setattr(workers, "fill_custom_template", lambda *args: (Path(template).read_bytes(), None, None, None))
        monkeypatch.setattr(workers, "upload_and_validate", lambda *args, **kw: pytest.fail("Invalid output was published"))
        result = await session.generate_document()
        assert not result["success"] and len(result["missing_field_ids"]) == 2
        snapshot = copy.deepcopy(session.doc)
        with pytest.raises(ValueError):
            session.apply_user_values({"T0_R0_C1": "伪造姓名"})
        assert session.doc == snapshot
    asyncio.run(run())


def test_rejected_new_source_invalidates_old_artifact_and_blocks_delivery():
    from runtime import PROJECT_ROOT
    template, paths = fixture_files()
    model = SmokeModel(template_path=template, file_paths=paths)
    async def run():
        session = DocumentSession({}, model, {}, "按材料生成", {"model_calls": 0, "usage_complete": True})
        await session.inspect_template(template)
        await session.extract_documents(paths)
        await session.match_fields()
        session.commit_proposal()
        assert (await session.generate_document())["success"]
        result = await session.extract_documents([str(PROJECT_ROOT / ".env")])
        assert not result["success"] and session.doc["pending_sources"]
        assert not session.doc.get("artifact")
        assert not (await session.generate_document())["success"]
    asyncio.run(run())


def test_modes_are_explicit_and_namespace_checkpoints(monkeypatch):
    monkeypatch.delenv("AGENT_MODE", raising=False)
    assert GraphService().mode == "workflow"
    ctx = Context()
    assert GraphService("leader_worker")._config({"session_id": "same"}, ctx)["configurable"]["thread_id"] == "leader_worker:same"
    assert GraphService("workflow")._config({"session_id": "same"}, ctx)["configurable"]["thread_id"] == "same"
    with pytest.raises(ValueError):
        GraphService("unknown")


def test_supervisor_http_sse_and_download_keep_public_contract(monkeypatch):
    import main
    from fastapi.testclient import TestClient
    template, paths = fixture_files()
    monkeypatch.setattr(main, "service", service_for(SmokeModel(template_path=template, file_paths=paths)))
    client = TestClient(main.app)
    result = client.post("/v1/chat/completions", json={"session_id": "http-leader", "stream": True,
        "messages": [{"role": "user", "content": "根据材料生成"}]})
    assert result.status_code == 200 and "data: [DONE]" in result.text
    chunks = [json.loads(line[6:]) for line in result.text.splitlines() if line.startswith("data: {")]
    assert not any("error" in chunk for chunk in chunks)
    outputs = [json.loads(delta["content"]) for chunk in chunks for choice in chunk.get("choices", [])
               if (delta := choice.get("delta", {})).get("role") == "tool" and delta["name"] == "generate_document"]
    assert outputs[-1]["success"]
    assert client.get(outputs[-1]["download_url"]).status_code == 200


def test_cancellation_cleans_up_supervisor_and_session_lock():
    class BlockingModel(SmokeModel):
        async def _agenerate(self, *args, **kwargs):
            await asyncio.sleep(100)

    template, paths = fixture_files()
    service = service_for(BlockingModel(template_path=template, file_paths=paths))
    async def run():
        ctx = Context()
        stream = service.stream_sse({"text": "等待", "session_id": "cancel-leader"}, ctx)
        task = asyncio.create_task(anext(stream))
        await asyncio.sleep(0.03)
        assert service.cancel_run(ctx.run_id)["status"] == "success"
        await asyncio.wait_for(task, timeout=2)
        await stream.aclose()
        assert not service.running_tasks and not service._active_sessions
    asyncio.run(run())


def test_resolving_one_conflict_requires_rematching_remaining_facts():
    template, paths = fixture_files("姓名：李四\n成绩：0")
    model = SmokeModel(template_path=template, file_paths=paths)

    async def run():
        session = DocumentSession({}, model, {}, "按材料生成", {"model_calls": 0, "usage_complete": True})
        await session.inspect_template(template)
        await session.extract_documents(paths)
        await session.match_fields()
        session.user_text = "姓名确认使用李四，请生成"
        session.apply_user_values({"T0_R0_C1": "李四"})
        assert not session.validate_document()["success"]
        assert not (await session.generate_document())["success"]
        await session.match_fields()
        session.commit_proposal()
        result = await session.generate_document()
        assert result["success"]
        output = Document(result["local_path"])
        assert output.tables[0].cell(0, 1).text == "李四"
        assert output.tables[0].cell(1, 1).text == "0"
    asyncio.run(run())


def test_explicit_numeric_override_cannot_use_a_substring():
    template, paths = fixture_files()

    async def run():
        session = DocumentSession({}, None, {}, "成绩改为100", {})
        await session.inspect_template(template)
        with pytest.raises(ValueError):
            session.apply_user_values({"T0_R1_C1": 0})
        assert not session.doc["values"]
        session.apply_user_values({"T0_R1_C1": 100})
        assert session.doc["values"]["T0_R1_C1"] == 100
    asyncio.run(run())


def test_worker_retry_receives_safe_validation_feedback():
    class FeedbackModel(SmokeModel):
        payloads: list[dict] = Field(default_factory=list)

        def _generate(self, messages, stop=None, run_manager=None, **kwargs):
            if messages[0].content.startswith("你是材料提取Worker"):
                self.payloads.append(json.loads(messages[-1].content))
                if len(self.payloads) == 1:
                    return ChatResult(generations=[ChatGeneration(message=AIMessage(content="invalid-json-sentinel"))])
            return super()._generate(messages, stop=stop, run_manager=run_manager, **kwargs)

    template, paths = fixture_files()
    model = FeedbackModel(template_path=template, file_paths=paths)
    result = asyncio.run(service_for(model).run({"text": "按材料生成", "session_id": "format-retry"}))
    assert result["delivered"]
    assert len(model.payloads) == 2
    feedback = model.payloads[1]["previous_error"]
    assert "JSON" in feedback and "invalid-json-sentinel" not in feedback
    assert not result["document"]["source_errors"]


def test_provider_value_error_does_not_expose_raw_error_to_leader():
    class FailingModel(SmokeModel):
        def _generate(self, messages, stop=None, run_manager=None, **kwargs):
            if messages[0].content.startswith("你是字段匹配Worker"):
                raise ValueError("provider-error-secret-sentinel")
            return super()._generate(messages, stop=stop, run_manager=run_manager, **kwargs)

    template, paths = fixture_files()
    result = asyncio.run(service_for(FailingModel(template_path=template, file_paths=paths)).run(
        {"text": "按材料生成", "session_id": "provider-error"}))
    assert not result["delivered"]
    assert "provider-error-secret-sentinel" not in "\n".join(str(m.content) for m in result["messages"])
