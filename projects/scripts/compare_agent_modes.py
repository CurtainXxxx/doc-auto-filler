"""Opt-in live A/B smoke checks on repository fixtures; saves evidence in data/.

Run: uv run python scripts/compare_agent_modes.py --live
Calls the configured paid provider. Never prints credentials or provider errors.
Without --live (or without model configuration), no network model calls occur.
"""
import argparse
import asyncio
import io
import json
import os
import re
import sys
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from docx import Document
from langchain_core.callbacks import BaseCallbackHandler
from langchain_core.messages import ToolMessage
from runtime import DATA_DIR, PROJECT_ROOT
from service import GraphService


class Meter(BaseCallbackHandler):
    def __init__(self):
        self.calls = 0
        self.input_tokens = 0
        self.output_tokens = 0
        self.usage_complete = True

    def on_chat_model_start(self, *args, **kwargs):
        self.calls += 1

    def on_llm_end(self, response, **kwargs):
        message = response.generations[0][0].message
        usage = getattr(message, "usage_metadata", None)
        if usage:
            self.input_tokens += usage.get("input_tokens", 0)
            self.output_tokens += usage.get("output_tokens", 0)
        else:
            self.usage_complete = False

    def on_llm_error(self, *args, **kwargs):
        self.usage_complete = False

    def report(self):
        return {"model_calls": self.calls, "input_tokens": self.input_tokens if self.usage_complete else None,
                "output_tokens": self.output_tokens if self.usage_complete else None,
                "usage_complete": self.usage_complete}


class MeasuredService(GraphService):
    def __init__(self, mode, meter):
        super().__init__(mode)
        self.meter = meter

    def _config(self, payload, ctx):
        config = super()._config(payload, ctx)
        config["callbacks"] = [self.meter]
        return config


def artifact_from(result):
    if "delivered" in result:
        return result["document"].get("artifact") if result["delivered"] else None
    for message in reversed(result.get("messages", [])):
        if isinstance(message, ToolMessage):
            try:
                data = json.loads(message.content)
                if data.get("success") and data.get("local_path") and data.get("download_url"):
                    return data
            except (ValueError, TypeError):
                pass
    return None


def cases():
    template = str(PROJECT_ROOT / "assets" / "templates" / "教务数据申请表.docx")
    knowledge = next((PROJECT_ROOT / "assets" / "knowledge" / "教务数据申请").glob("*简单*.txt"))
    from document_access import save_upload
    conflicts = [save_upload(f"申请时间：{date}\n申请人姓名：李华".encode(), ".txt")
                 for date in ("2026年10月20日", "2026年10月22日")]
    return [
        {"name": "single_document", "template": template, "files": [str(knowledge)],
         "expected": {(0, 1): "教务处教学运行科", (0, 3): "2024年3月10日", (1, 1): "李华",
                      (1, 3): "20190567", (2, 1): "13812345678", (2, 3): "lihua@lnnu.edu.cn"},
         "clarification": False},
        {"name": "conflicting_dates", "template": template, "files": conflicts,
         "expected": {}, "clarification": True},
    ]


async def run_case(case, mode, output):
    meter = Meter()
    service = MeasuredService(mode, meter)
    session = f"compare-{mode}-{uuid.uuid4().hex}"
    prompt = (f"模板：{case['template']}。材料文件：{json.dumps(case['files'], ensure_ascii=False)}。"
              f"请读取全部材料并填写模板，表单会话ID使用{session}。不要编造或代签。"
              "申请数据描述、数据用途、申请人签名和负责人签名均留空。"
              "若同一字段在材料中出现不同值，先向我澄清，禁止擅自选择或生成；否则请生成Word并提供下载。")
    record = {"mode": mode, "case": case["name"], "live": True, "status": "failed"}
    started = time.perf_counter()
    try:
        async with asyncio.timeout(180):
            result = await service.run({"session_id": session, "text": prompt})
        artifact = artifact_from(result)
        if case["clarification"]:
            answer = next((m.content for m in reversed(result["messages"]) if m.type == "ai" and isinstance(m.content, str)), "")
            clarified = any(word in answer for word in ("冲突", "不一致", "不同日期")) and any(
                word in answer for word in ("确认", "澄清", "哪", "选择"))
            record.update(status="passed" if artifact is None and clarified else "failed",
                          clarification_without_generation=artifact is None and clarified)
        elif artifact:
            data = Path(artifact["local_path"]).read_bytes()
            doc = Document(io.BytesIO(data))
            from tools.filling.base import get_unique_cells
            checks = [{"row": row, "col": col, "expected": value,
                       "actual": get_unique_cells(doc.tables[0].rows[row])[col].text.strip()}
                      for (row, col), value in case["expected"].items()]
            for check in checks:
                check["passed"] = check["actual"] == check["expected"]
            source = Document(case["template"])
            placeholders = all(get_unique_cells(doc.tables[0].rows[row])[1].text ==
                               get_unique_cells(source.tables[0].rows[row])[1].text for row in (3, 4))
            def signatures(document):
                text = "\n".join(cell.text for row in document.tables[0].rows for cell in get_unique_cells(row))
                return re.findall(r"(?:申请人|负责人)签名[^\n]*", text)
            blanks = placeholders and signatures(doc) == signatures(source)
            target = output / f"{case['name']}-{mode}.docx"
            target.write_bytes(data)
            record.update(status="passed" if all(c["passed"] for c in checks) and blanks else "failed",
                          checks=checks, requested_blank_fields_preserved=blanks, file=str(target))
        record["trace"] = result.get("trace")
        record["budget_exhausted"] = result.get("budget_exhausted", False)
    except Exception as exc:
        record["error_type"] = type(exc).__name__
    record.update(seconds=round(time.perf_counter() - started, 3), **meter.report())
    print(json.dumps(record, ensure_ascii=False), flush=True)
    return record


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--live", action="store_true", help="Actually call the configured provider")
    parser.add_argument("--mode", choices=("both", "workflow", "leader_worker"), default="both")
    args = parser.parse_args()
    output = DATA_DIR / "leader-worker-compare" / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    output.mkdir(parents=True, exist_ok=True)
    configured = all(os.getenv(k) for k in ("EXTERNAL_LLM_BASE_URL", "EXTERNAL_LLM_MODEL", "EXTERNAL_LLM_API_KEY"))
    if not args.live or not configured:
        report = {"status": "not_run", "reason": "explicit --live and complete local model configuration required"}
    else:
        async def run():
            records = []
            for case in cases():
                for mode in (("workflow", "leader_worker") if args.mode == "both" else (args.mode,)):
                    records.append(await run_case(case, mode, output))
            return records
        report = {"status": "completed", "results": asyncio.run(run()),
                  "limits": "one trial per case; not a statistical accuracy or latency benchmark"}
    path = output / "results.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Comparison status: {report['status']}; evidence: {path}")
    return 2 if report["status"] == "not_run" else int(any(r["status"] != "passed" for r in report["results"]))


if __name__ == "__main__":
    raise SystemExit(main())
