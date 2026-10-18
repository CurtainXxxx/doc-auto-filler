"""Exercise an independently running server over TCP with shipped Word templates.

Run from projects/: uv run python scripts/live_smoke.py
This checks deterministic document operations; it does not substitute for AI tests.
Only synthetic fixture data is used. Downloaded files and JSON evidence stay in data/.
"""
import argparse
import hashlib
import io
import json
import time
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import httpx
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]


def cases():
    return [
        ("教务数据申请表", "教务数据申请", {
            "数据申请部门": "教务处教学运行科", "申请时间": "2024年3月10日",
            "申请人姓名": "李华", "工资号": "20190567", "联系电话": "13812345678",
            "电子邮箱": "lihua@lnnu.edu.cn", "申请数据描述": "2024级课程选课统计（验收样例）",
            "数据用途": "教学运行管理（验收样例）",
        }, []),
        ("教材建设申报书", "教材建设申报", {
            "P9": "人工智能导论——从原理到实践", "P10": "周建国",
            "P11": "计算机科学与技术学院", "P12": "13987654321", "P13": "2026年10月18日",
            "T0_R0_C1": "周建国", "专业技术职务": "教授", "电子邮箱": "zhoujg@ustc.edu.cn",
            "T1_R0_C1": "人工智能导论——从原理到实践", "适用对象": "计算机专业本科生",
        }, []),
        ("考场记录表", "考场记录表", {
            "考试课程": "数据结构", "考试时间": "2024年1月15日 上午 9:00-11:00",
            "考生所在院系": "计算机与智能教育学院", "考试教室": "第二教学楼2203",
            "应到人数": 45, "实到人数": 43, "考试方式": "闭卷",
            "主讲教师签名": "张明", "班主任签名": "刘伟",
            "T0_G0": [[f"TEST2026{i:04d}", f"验收考生{i}", "测试学院", "√", "", "", "", "", "已核实"]
                       for i in range(1, 8)],
            "T0_G1": [["验收用扩展行，非真实考生资料。"]],
        }, [(0, 6, {"闭卷"})]),
        ("评价报告模板", "评价报告", {
            "课程名称": "数据结构", "开课时间": "2023-2024学年第二学期",
            "教学班级": "计科2201-2202", "评价责任人": "王明", "参与人": "李华、赵芳",
            "课程总结": "验收样例：课程目标基本达成。", "改进措施": "验收样例：增加算法实践。",
        }, []),
        ("试卷分析模板", "试卷分析", {
            "课程名称": "数据结构与算法", "学时数": 64, "开课单位": "计算机科学与技术学院",
            "教师姓名": "王明", "课程性质": "必修", "命题形式": "教师组题",
            "考试形式": "闭卷", "教考分离": "否", "学生班级": "计算机2301班", "分数分布_<60": 3,
            "分数分布_60-69": 7, "分数分布_70-79": 15, "分数分布_80-89": 12,
            "分数分布_90-100": 9, "平均分": 76.8, "最高分": 97, "最低分": 38, "标准差": 14.2,
            "学生班级_作弊": 0,
        }, [(0, 3, {"必修"}), (0, 4, {"教师组题"}), (0, 5, {"闭卷", "否"})]),
        ("关联矩阵模板", "关联矩阵", {
            "T2_R2_C1": 10, "T2_R2_C2": "数组与链表", "T2_R2_C3": "√",
            "T2_R2_C9": "选择题", "T2_R2_C10": 0.8,
            "T0_G0": [["工程知识", "基础知识运用", "掌握数据结构原理"]],
            "T1_G1": [["课程目标1", "1、2、3"]],
        }, []),
    ]


def checked_json(response):
    response.raise_for_status()
    body = response.json()
    if not body.get("success"):
        raise AssertionError(body.get("message") or "API returned success=false")
    return body


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:5000")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data" / "live-http-results")
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    evidence = {"started_at": datetime.now(timezone.utc).isoformat(), "transport": "real TCP HTTP",
                "mocked_components": [], "cases": [],
                "not_verified": ["real LLM", "S3", "PostgreSQL", "PDF visual layout", "arbitrary templates"]}
    with httpx.Client(base_url=args.base_url, timeout=120, trust_env=False) as client:
        health = client.get("/health")
        health.raise_for_status()
        assert health.json()["status"] == "ok"
        front = client.get("/web/")
        front.raise_for_status()
        assert "<html" in front.text.lower()
        for name, knowledge_name, values, checkbox_rows in cases():
            started = time.monotonic()
            result = {"template": name, "status": "failed", "issues": []}
            try:
                template = ROOT / "assets" / "templates" / f"{name}.docx"
                original = template.read_bytes()
                source_hash = hashlib.sha256(original).hexdigest()
                knowledge = next((ROOT / "assets" / "knowledge" / knowledge_name).glob("*简单*.txt"))
                uploaded_knowledge = checked_json(client.post("/upload", files={"files": (knowledge.name, knowledge.read_bytes())}))
                result["knowledge_uploaded"] = bool(uploaded_knowledge.get("file_path"))
                uploaded = checked_json(client.post("/upload-template", files={"file": (template.name, original)}))
                preview = checked_json(client.get("/template-preview", params={"path": uploaded["template_path"]}))
                assert preview["field_map"] and preview["html"]
                generated = checked_json(client.post("/cofill-generate", json={
                    "template_path": uploaded["template_path"], "field_values": values}))
                response = client.get(generated["download_url"])
                response.raise_for_status()
                output = args.output_dir / f"{name}-实测.docx"
                output.write_bytes(response.content)
                doc = Document(io.BytesIO(response.content))
                with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                    assert archive.testzip() is None
                    for part in archive.namelist():
                        if part.endswith(".xml"):
                            etree.fromstring(archive.read(part))
                source_doc = Document(io.BytesIO(original))
                assert len(doc.tables) == len(source_doc.tables)
                assert [etree.tostring(s) for s in doc._element.iter(qn("w:sectPr"))] == [
                    etree.tostring(s) for s in source_doc._element.iter(qn("w:sectPr"))]
                text = " ".join(doc._element.itertext())
                expected = []
                for value in values.values():
                    if isinstance(value, list):
                        expected.extend(str(cell) for row in value for cell in row if cell != "")
                    else:
                        expected.append(str(value))
                compact_text = text.replace(" ", "").replace("\u3000", "")
                missing = sorted({v for v in expected if v not in text and v.replace(" ", "") not in compact_text})
                if missing:
                    result["issues"].append({"missing_values": missing})
                # Locate options independently from the source XML: each option
                # cell is followed by its (possibly merged) marking cell.
                for table_idx, row_idx, selected in checkbox_rows:
                    source_cells = source_doc.tables[table_idx].rows[row_idx]._tr.findall(qn("w:tc"))
                    output_cells = doc.tables[table_idx].rows[row_idx]._tr.findall(qn("w:tc"))
                    for idx, tc in enumerate(source_cells[:-1]):
                        option = "".join(t.text or "" for t in tc.iter(qn("w:t"))).strip().replace(" ", "").replace("\u3000", "")
                        if option not in {"开卷", "闭卷", "半开卷", "其他", "选修", "必修", "试题库", "试卷库", "教师组题", "是", "否"}:
                            continue
                        actual = "".join(t.text or "" for t in output_cells[idx + 1].iter(qn("w:t"))).strip()
                        if actual != ("√" if option in selected else ""):
                            result["issues"].append({"wrong_checkbox": [table_idx, row_idx, option], "actual": actual})
                checked_json(client.get("/generated-preview", params={"local_path": generated["local_path"]}))
                assert source_hash == hashlib.sha256(template.read_bytes()).hexdigest()
                result.update(status="passed" if not result["issues"] else "failed", output=str(output),
                              expected_value_checks=len(expected), bytes=len(response.content),
                              tables=len(doc.tables), source_sha256=source_hash,
                              generated_sha256=hashlib.sha256(response.content).hexdigest(),
                              uploaded_template=uploaded["template_path"],
                              uploaded_knowledge=uploaded_knowledge["file_path"],
                              generated_path=generated["local_path"])
            except Exception as exc:
                result["issues"].append({"exception_type": type(exc).__name__, "message": str(exc)})
            result["seconds"] = round(time.monotonic() - started, 3)
            evidence["cases"].append(result)
            print(json.dumps({k: result[k] for k in ("template", "status", "issues", "seconds")}, ensure_ascii=False))
    evidence["passed"] = sum(c["status"] == "passed" for c in evidence["cases"])
    evidence["failed"] = len(evidence["cases"]) - evidence["passed"]
    (args.output_dir / "results.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Passed {evidence['passed']}/6; evidence: {args.output_dir / 'results.json'}")
    return bool(evidence["failed"])


if __name__ == "__main__":
    raise SystemExit(main())
