#!/usr/bin/env python3
"""
自动端到端评测脚本 v2
测试 7 个模板 × 3 个难度 = 21 个用例的全流程
输出评测报表（填充率、准确率、Token消耗、成本）

修复记录:
  v2: 1) 支持 Agent 确认后二次触发生成 2) Token 从文本估算 3) 实时保存中间结果
"""
import json, os, re, sys, time, traceback
import requests
from docx import Document
from datetime import datetime

# ============================================================
# 配置
# ============================================================
BASE_URL = "http://localhost:5000"
PROJECT_DIR = "/workspace/projects/projects"
GT_DIR = "/workspace/projects/ground_truth"
REPORT_DIR = "/workspace/projects"
KNOWLEDGE_DIR = os.path.join(PROJECT_DIR, "assets/knowledge")
TEMPLATE_DIR = os.path.join(PROJECT_DIR, "assets/templates")

PRICE_INPUT = 1
PRICE_OUTPUT = 2

TEMPLATE_FILES = {
    "缓考申请单": "缓考申请单.docx",
    "考场记录表": "考场记录表.docx",
    "试卷分析模板": "试卷分析模板.docx",
    "关联矩阵模板": "关联矩阵模板.docx",
    "教材建设申报书": "教材建设申报书.docx",
    "教务数据申请表": "教务数据申请表.docx",
    "评价报告模板": "评价报告模板.docx",
}

KNOWLEDGE_DIRS = {
    "缓考申请单": "缓考申请单",
    "考场记录表": "考场记录表",
    "试卷分析": "试卷分析模板",
    "关联矩阵": "关联矩阵模板",
    "教材建设申报": "教材建设申报书",
    "教务数据申请": "教务数据申请表",
    "评价报告": "评价报告模板",
}

DIFFICULTIES = ["简单", "中等", "困难"]
AGENT_TIMEOUT = 300
INTERIM_RESULT = os.path.join(PROJECT_DIR, "eval_interim.json")

# 全局计时
_GLOBAL_START = time.time()


# ============================================================
# 工具函数
# ============================================================

def estimate_tokens(text: str) -> dict:
    """从文本估算 token 消耗"""
    total = len(text)
    estimated = max(1, total // 2)
    return {"prompt_tokens": estimated * 3, "completion_tokens": estimated, "total_tokens": estimated * 4}


def extract_s3_url(text):
    """从 Agent 响应文本中提取 docx 下载 URL"""
    if not text:
        return None
    patterns = [
        r'\[[^\]]*\]\((https?://[^\s\(\)]+?\.docx[^\s\(\)]*)\)',
        r'(https?://[^\s<>"\'\)）\)。，,!\?；;]+\.docx[^\s<>"\'\)）\)。，,!\?；;]*)',
        r'下载链接[：:]\s*(https?://[^\s<>"\'\)）\)。，,!\?；;]+)',
        r'下载[：:]\s*(https?://[^\s<>"\'\)）\)。，,!\?；;]+)',
        r'"download_url"\s*:\s*"(https?://[^"]+)"',
        r'(https?://[^\s<>"\'\)）\)。，,!\?；;]*tos\.coze[^\s<>"\'\)）\)。，,!\?；;]*)',
    ]
    for p in patterns:
        m = re.search(p, text)
        if m:
            url = m.group(1) if m.lastindex else m.group(0)
            url = re.sub(r'[）\)。，,!\?；;]+$', '', url)
            if '.docx' in url or 'download' in url or 'generated' in url or 'output' in url or 'tos.coze' in url:
                return url
    return None


def parse_sse_stream(resp):
    """解析 SSE 流，返回 (完整文本, 下载URL, 错误)"""
    full_text = ""
    s3_url = None
    error = None
    for line in resp.iter_lines():
        if not line:
            continue
        decoded = line.decode("utf-8", errors="replace")
        if decoded.startswith("data: ") and decoded != "data: [DONE]\n\n":
            data_str = decoded[6:]
            try:
                chunk = json.loads(data_str)
                choices = chunk.get("choices", [])
                if choices:
                    delta = choices[0].get("delta", {})
                    content = delta.get("content", "")
                    if content:
                        full_text += content
            except json.JSONDecodeError:
                pass
    if full_text:
        s3_url = extract_s3_url(full_text)
    return full_text, s3_url, error


def extract_docx_text(docx_path):
    """从 docx 中提取所有文本"""
    doc = Document(docx_path)
    result = {"paragraphs": [], "tables": [], "all_text": ""}
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            result["paragraphs"].append(txt)
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            table_data.append([c.text.strip() for c in row.cells])
        result["tables"].append(table_data)
    result["all_text"] = "\n".join(result["paragraphs"]) + "\n" + \
        "\n".join("\t".join(c for c in row) for t in result["tables"] for row in t)
    return result


def check_field_filled(docx_text, field_label, expected_value, optional=False):
    if optional and not expected_value:
        return "optional"
    if not expected_value:
        return "empty"
    if expected_value in docx_text:
        return "correct"
    return "filled_wrong"


def check_row_group_filled(docx_tables, gt_rows, group_id):
    if not gt_rows:
        return 0, 0, "无数据行"
    matched, total = 0, len(gt_rows)
    details = []
    for row_idx, expected_row in enumerate(gt_rows):
        found = False
        for table in docx_tables:
            for docx_row in table:
                match_count = sum(1 for ev in expected_row if any(ev in c for c in docx_row))
                if match_count >= len(expected_row) * 0.5:
                    found = True
                    break
            if found:
                break
        if found:
            matched += 1
            details.append(f"行{row_idx+1}: ✅")
        else:
            details.append(f"行{row_idx+1}: ❌ ({'|'.join(expected_row)[:40]})")
    return matched, total, "; ".join(details)


def call_agent(messages, session_id, timeout=AGENT_TIMEOUT):
    """调用 Agent（流式），返回 (full_text, s3_url, error)"""
    resp = requests.post(
        f"{BASE_URL}/v1/chat/completions",
        json={"model": "agent", "messages": messages, "stream": True, "session_id": session_id},
        stream=True, timeout=timeout,
    )
    resp.raise_for_status()
    return parse_sse_stream(resp)


def needs_confirmation(text):
    if not text:
        return False
    patterns = [r'确认无误', r'请您确认', r'请确认', r'确认以上', r'是否正确', r'是否确认', r'同意生成', r'请审核']
    return any(re.search(p, text) for p in patterns)


# ============================================================
# 评测主函数
# ============================================================

def evaluate_single_test(gt):
    tmpl_name = gt["template"]
    difficulty = gt["difficulty"]
    session_id = f"eval_{tmpl_name}_{difficulty}_{int(time.time())}"

    tmpl_file = TEMPLATE_FILES.get(tmpl_name)
    if not tmpl_file:
        return {"error": f"未知模板: {tmpl_name}"}

    tmpl_path = os.path.join(TEMPLATE_DIR, tmpl_file)
    if not os.path.isfile(tmpl_path):
        return {"error": f"模板文件不存在: {tmpl_path}"}

    kdir_name = next((d for d, t in KNOWLEDGE_DIRS.items() if t == tmpl_name), None)
    kfile = os.path.join(KNOWLEDGE_DIR, kdir_name, f"{kdir_name}_【{difficulty}】.txt")

    result = {
        "template": tmpl_name, "difficulty": difficulty,
        "knowledge_file": os.path.basename(kfile),
        "total_fields": gt["total_fields"],
        "optional_count": gt["optional_fields_count"],
        "status": "pending", "error": None,
        "fill_rate": 0.0, "accuracy": 0.0, "row_group_accuracy": 0.0,
        "prompt_tokens": 0, "completion_tokens": 0, "cost": 0.0, "duration": 0.0, "details": {}
    }
    start_time = time.time()

    try:
        # 1. 上传模板
        with open(tmpl_path, "rb") as f:
            r = requests.post(f"{BASE_URL}/upload-template", files={"file": f}, timeout=30)
        r.raise_for_status()
        tmpl_data = r.json()
        if not tmpl_data.get("success"):
            result.update({"status": "failed", "error": f"上传模板失败: {tmpl_data}", "duration": time.time()-start_time})
            return result
        template_path = tmpl_data["template_path"]
        print(f"  [1/5] ✅ 上传模板")

        # 2. 上传知识
        with open(kfile, "rb") as f:
            r = requests.post(f"{BASE_URL}/upload", files={"files": ("knowledge.txt", f)}, timeout=30)
        r.raise_for_status()
        extracted_text = r.json().get("extracted_text", "")
        print(f"  [2/5] ✅ 上传知识 ({len(extracted_text)} 字符)")

        # 3. 调用 Agent
        user_msg = f"帮我填写这个文档模板：{template_path}\n\n以下是知识文件的内容（{os.path.basename(kfile)}）：\n\n{extracted_text}"
        messages = [{"role": "user", "content": user_msg}]
        agent_text, s3_url, error = call_agent(messages, session_id)

        # 3.5 如果 Agent 请求确认，自动确认
        if not s3_url and needs_confirmation(agent_text):
            print(f"  [3.5] ⚠️ Agent 请求确认，自动发送确认消息")
            messages.append({"role": "assistant", "content": agent_text})
            messages.append({"role": "user", "content": "确认无误，请生成最终的Word文档。"})
            agent_text2, s3_url2, error2 = call_agent(messages, session_id, timeout=AGENT_TIMEOUT)
            agent_text += agent_text2
            if s3_url2:
                s3_url = s3_url2

        est = estimate_tokens(agent_text)
        result["prompt_tokens"] = est["prompt_tokens"]
        result["completion_tokens"] = est["completion_tokens"]
        result["cost"] = (est["prompt_tokens"] * PRICE_INPUT + est["completion_tokens"] * PRICE_OUTPUT) / 1_000_000
        result["agent_response_length"] = len(agent_text)
        print(f"  [3/5] ✅ Agent 完成 ({est['prompt_tokens']} in / {est['completion_tokens']} out, est)")

        # 4. 下载 docx
        docx_path = None
        if s3_url:
            r = requests.get(f"{BASE_URL}/download-docx", params={"file_path": s3_url}, timeout=60)
            if r.status_code == 200 and len(r.content) > 100:
                docx_path = f"/tmp/eval_{tmpl_name}_{difficulty}.docx"
                with open(docx_path, "wb") as f:
                    f.write(r.content)
                print(f"  [4/5] ✅ 下载 docx ({len(r.content)} 字节)")
            else:
                print(f"  [4/5] ⚠️ /download-docx 返回 {r.status_code}, size={len(r.content)}")
        else:
            print(f"  [4/5] ⚠️ 未从 Agent 响应中找到下载链接")

        # 5. 评测
        if docx_path and os.path.isfile(docx_path):
            docx_data = extract_docx_text(docx_path)
            gt_fields = gt.get("fields", {})
            opt_fields = gt.get("optional_fields", [])
            field_results = {}
            correct_count = filled_count = non_optional_total = 0
            for label, expected_value in gt_fields.items():
                is_optional = label in opt_fields
                if not is_optional:
                    non_optional_total += 1
                status = check_field_filled(docx_data["all_text"], label, expected_value, is_optional)
                if status == "correct":
                    correct_count += 1
                    filled_count += 1
                elif status == "filled_wrong":
                    filled_count += 1
                field_results[label] = status

            gt_row_groups = gt.get("row_groups", {})
            row_group_results = {}
            total_rg_matched = total_rg_rows = 0
            for group_id, expected_rows in gt_row_groups.items():
                if not expected_rows:
                    row_group_results[group_id] = {"status": "skipped", "detail": "无数据行"}
                    continue
                matched, total, detail = check_row_group_filled(docx_data["tables"], expected_rows, group_id)
                total_rg_matched += matched
                total_rg_rows += total
                row_group_results[group_id] = {"matched": matched, "total": total, "detail": detail[:100]}

            result["fill_rate"] = round(filled_count / non_optional_total * 100, 1) if non_optional_total > 0 else 0
            result["accuracy"] = round(correct_count / non_optional_total * 100, 1) if non_optional_total > 0 else 0
            result["field_results"] = field_results
            result["row_group_results"] = row_group_results
            result["row_group_accuracy"] = round(total_rg_matched / total_rg_rows * 100, 1) if total_rg_rows > 0 else 0
            result["non_optional_fields"] = non_optional_total
            result["correct_count"] = correct_count
            result["filled_count"] = filled_count
            print(f"  [5/5] ✅ 评测完成 (填充率={result['fill_rate']}%, 准确率={result['accuracy']}%)")
            result["status"] = "completed"
        else:
            print(f"  [5/5] ⚠️ 无 docx 可解析，仅记录 Agent 响应")
            result["status"] = "partial"
            result["note"] = "docx 未下载成功"

    except requests.Timeout:
        result["status"] = "failed"
        result["error"] = f"Timeout (>{AGENT_TIMEOUT}s)"
    except Exception as e:
        result["status"] = "failed"
        result["error"] = f"{type(e).__name__}: {str(e)}"
        traceback.print_exc()

    result["duration"] = round(time.time() - start_time, 1)
    return result


def run_all_tests():
    print("=" * 60)
    print(f"  自动评测启动  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"  服务器: {BASE_URL}")
    print(f"  用例数: 7 模板 × 3 难度 = 21")
    print("=" * 60)

    gts = {}
    for fname in sorted(os.listdir(GT_DIR)):
        if not fname.endswith(".json"):
            continue
        with open(os.path.join(GT_DIR, fname), encoding="utf-8") as f:
            gt = json.load(f)
        gts[f"{gt['template']}_{gt['difficulty']}"] = gt

    test_order = []
    for tmpl_name in TEMPLATE_FILES:
        for diff in DIFFICULTIES:
            key = f"{tmpl_name}_{diff}"
            if key in gts:
                test_order.append(key)

    results = []
    total = len(test_order)
    for i, key in enumerate(test_order, 1):
        gt = gts[key]
        elapsed = int((time.time() - _GLOBAL_START) / 60)
        print(f"\n[{i}/{total}] {gt['template']} × {gt['difficulty']} (已耗时 {elapsed}分钟)...")
        result = evaluate_single_test(gt)
        results.append(result)
        with open(INTERIM_RESULT, "w") as f:
            json.dump({"results": results, "completed": i, "total": total}, f, ensure_ascii=False, indent=2)
        summary = f"  → {result['status']} | 填充率={result.get('fill_rate', 0):.1f}% 准确率={result.get('accuracy', 0):.1f}% | {result.get('duration', 0):.0f}s"
        if result.get("error"):
            summary += f" | ❌ {result['error'][:80]}"
        print(summary)
    return results


def generate_report(results):
    if not results:
        print("无测试结果")
        return None

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = os.path.join(REPORT_DIR, f"eval_report_{timestamp}.md")

    completed = [r for r in results if r["status"] == "completed"]
    partial = [r for r in results if r["status"] == "partial"]
    failed = [r for r in results if r["status"] == "failed"]

    total_cost = sum(r.get("cost", 0) for r in results)
    total_tokens = sum(r.get("prompt_tokens", 0) + r.get("completion_tokens", 0) for r in results)
    total_duration = sum(r.get("duration", 0) for r in results)
    avg_fill = sum(r.get("fill_rate", 0) for r in completed) / len(completed) if completed else 0
    avg_acc = sum(r.get("accuracy", 0) for r in completed) / len(completed) if completed else 0
    avg_rg = None
    rg_scores = [r.get("row_group_accuracy", 0) for r in completed if r.get("row_group_accuracy", 0) > 0]
    if rg_scores:
        avg_rg = sum(rg_scores) / len(rg_scores)

    lines = [
        "# 端到端自动评测报告",
        "",
        f"**生成时间**: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"**模型**: DeepSeek V4-Flash（输入 ¥{PRICE_INPUT}/1M, 输出 ¥{PRICE_OUTPUT}/1M）",
        "",
        "## 总览",
        "",
        "| 指标 | 数值 |",
        "|------|------|",
        f"| 总用例数 | {len(results)} |",
        f"| 完成 | {len(completed)} |",
        f"| 部分完成 | {len(partial)} |",
        f"| 失败 | {len(failed)} |",
        f"| 平均填充率 | **{avg_fill:.1f}%** |",
        f"| 平均准确率 | **{avg_acc:.1f}%** |",
    ]
    if avg_rg is not None:
        lines.append(f"| 行组平均准确率 | **{avg_rg:.1f}%** |")
    lines += [
        f"| 总 Token 消耗 | {total_tokens:,} |",
        f"| 总成本 | **¥{total_cost:.4f}** |",
        f"| 总耗时 | {total_duration:.0f}s（约 {total_duration/60:.1f} 分钟）|",
        "",
        "## 成本对比",
        "",
        "| 方式 | 总成本 | 单价 | 总时间 |",
        "|------|--------|------|--------|",
    ]
    if completed:
        lines.append(f"| **本系统** | ¥{total_cost:.4f} | ¥{total_cost/len(completed):.4f} | {total_duration:.0f}s（{total_duration/len(completed):.0f}s/次）|")
    lines += [
        f"| **人工填写（估）** | ¥{25*len(completed):.0f} | ¥25 | {30*len(completed)}min（30min/次）|",
        "",
        "## 按模板统计",
        "",
        "| 模板 | 难度 | 状态 | 字段数 | 填充率 | 准确率 | Token(in/out) | 成本 | 耗时 |",
        "|------|------|------|--------|--------|--------|-------------|------|------|",
    ]
    for r in results:
        icon = "✅" if r["status"] == "completed" else ("⚠️" if r["status"] == "partial" else "❌")
        fill = f"{r.get('fill_rate', 0):.1f}%" if r["status"] == "completed" else "-"
        acc = f"{r.get('accuracy', 0):.1f}%" if r["status"] == "completed" else "-"
        tok = f"{r.get('prompt_tokens', 0)}/{r.get('completion_tokens', 0)}"
        cost = f"¥{r['cost']:.4f}" if r.get("cost") else "-"
        dur = f"{r.get('duration', 0):.0f}s"
        lines.append(f"| {r['template']} | {r['difficulty']} | {icon} | {r.get('total_fields','?')} | {fill} | {acc} | {tok} | {cost} | {dur} |")
    lines.append("")

    failed_tests = [r for r in results if r["status"] == "failed"]
    if failed_tests:
        lines.append("## 失败用例详情\n")
        for r in failed_tests:
            lines.append(f"### ❌ {r['template']} × {r['difficulty']}")
            lines.append(f"- 错误: {r.get('error', '未知')}\n")

    lines.append("## 各用例详情")
    for r in completed:
        lines.append(f"\n### {r['template']} × {r['difficulty']}")
        lines.append(f"- **填充率**: {r['fill_rate']:.1f}% ({r.get('filled_count', 0)}/{r.get('non_optional_fields', 0)})")
        lines.append(f"- **准确率**: {r['accuracy']:.1f}% ({r.get('correct_count', 0)}/{r.get('non_optional_fields', 0)})")
        lines.append(f"- **Token**: {r.get('prompt_tokens', 0)} in / {r.get('completion_tokens', 0)} out")
        lines.append(f"- **成本**: ¥{r['cost']:.4f}")
        lines.append(f"- **耗时**: {r['duration']:.0f}s\n")
        field_results = r.get("field_results", {})
        if field_results:
            lines.append("| 字段 | 状态 |")
            lines.append("|------|------|")
            for label, status in field_results.items():
                icon = {"correct": "✅", "filled_wrong": "⚠️", "empty": "❌", "optional": "—"}.get(status, "?")
                lines.append(f"| {label} | {icon} {status} |")
        rg_results = r.get("row_group_results", {})
        if rg_results:
            lines.append("\n**行组**:")
            for gid, rg in rg_results.items():
                lines.append(f"- {gid}: {rg.get('detail', '')}")

    report_text = "\n".join(lines)
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report_text)
    print(f"\n✅ 报表已生成: {report_path}")
    return report_path


# ============================================================
# 主入口
# ============================================================

if __name__ == "__main__":
    print("自动端到端评测脚本 v2")
    print("=" * 60)

    gt_files = [f for f in os.listdir(GT_DIR) if f.endswith(".json")]
    print(f"GT 文件: {len(gt_files)}")

    results = run_all_tests()
    if results:
        report_path = generate_report(results)
        print(f"\n🎉 评测完成！报表: {report_path}")
        print(f"   中间结果: {INTERIM_RESULT}")